"""Tests de non-régression issus de l'audit de la version 2.5."""

import io
import math
from datetime import date

import openpyxl
import pytest

from renovation_gantt.budget import Budget, Financeur, budget_vers_excel
from renovation_gantt.export import Scenario, planning_vers_excel, planning_vers_ics
from renovation_gantt.phases import EtatProjet
from renovation_gantt.saisie import borner

DEBUT = date(2026, 10, 14)


@pytest.mark.parametrize(
    ("valeur", "mini", "maxi", "attendu"),
    [
        (10, 1, 156, (10, False)),
        (10**7, 1, 156, (156, True)),
        (-5, 1, 104, (1, True)),
        (99, 0, 4, (4, True)),
        (math.nan, 0, 10, (0, True)),
        ("abc", 0.0, 100.0, (0.0, True)),
        (12.5, 0.0, 100.0, (12.5, False)),
    ],
)
def test_borner(valeur, mini, maxi, attendu):
    assert borner(valeur, mini, maxi) == attendu


def test_aucune_formule_dans_les_exports_excel():
    planning = Scenario(EtatProjet.MOE_SELECTIONNEE, DEBUT, 12).planning()
    budget = Budget(travaux_ht=100_000, financeurs=[Financeur('=HYPERLINK("http://exemple.org";"clic")', 10)])
    for contenu in (budget_vers_excel(budget, planning), planning_vers_excel(planning)):
        classeur = openpyxl.load_workbook(io.BytesIO(contenu))
        for feuille in classeur.worksheets:
            for ligne in feuille.iter_rows():
                assert all(c.data_type != "f" for c in ligne)
    valeur = openpyxl.load_workbook(io.BytesIO(budget_vers_excel(budget, planning)))["Plan de financement"]["A2"].value
    assert valeur.startswith("=HYPERLINK")  # conservé, mais comme simple texte


def test_lignes_ics_pliees_a_75_octets():
    planning = Scenario(EtatProjet.SANS_AUDIT, DEBUT, 12).planning()
    nom = "Rénovation énergétique du groupe scolaire Jean Moulin, de son gymnase et de la cantine municipale"
    texte = planning_vers_ics(planning, nom).decode()
    lignes = texte.split("\r\n")
    assert all(len(ligne.encode()) <= 75 for ligne in lignes)
    deplie = texte.replace("\r\n ", "")
    assert f"X-WR-CALNAME:{nom.replace(',', chr(92) + ',')}" in deplie


def test_budget_lu_dans_un_fichier_est_assaini():
    budget = Budget.depuis_dict(
        {"financeurs": [{"nom": "A", "taux": 250, "montant": -10}, {"nom": "B", "taux": "NaN"}]}
    )
    assert [(f.taux, f.montant) for f in budget.financeurs] == [(100.0, 0.0), (0.0, 0.0)]


def test_cout_du_budget_calcule_une_seule_fois():
    budget = Budget(travaux_ht=500_000)
    assert budget.total_ht == budget.total_ht
    assert budget._couts is budget._couts
    copie = budget.couts()
    copie.loc[0, "HT"] = 0  # modifier la copie renvoyée ne change pas le budget
    assert budget.couts().loc[0, "HT"] == 500_000
